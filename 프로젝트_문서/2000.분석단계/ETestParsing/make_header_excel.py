#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
헤더 분해 결과를 Excel 로 만든다 (CLAUDE-054).
  시트 1 "헤더분해": 헤더 1개 = 1행, 맨 왼쪽 원본 헤더, 오른쪽에 분해 컬럼 + 적용 Rule
  시트 2 "SpecRule": spec_rules.csv 그대로

입력: out/item_spec_map.csv (spec_rule_tool.py apply 결과), spec_rules.csv
  python make_header_excel.py            -> out/header_split.xlsx
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAP_CSV, RULE_CSV, OUT_XLSX = "out/item_spec_map.csv", "spec_rules.csv", "out/header_split.xlsx"

# (시트 제목, item_spec_map 컬럼) - 첫 열은 원본 헤더
SPLIT_COLS = [
    ("원본 헤더", "RAW_HEADER"),
    ("TEST", "TEST"), ("COND_NAME", "COND_NAME"), ("COND_VALUE", "COND_VALUE"), ("COND_UNIT", "COND_UNIT"),
    ("STRUCT_TYPE", "STRUCT_TYPE"), ("ITEM_ID", "ITEM_ID"), ("PIN", "PIN"),
    ("X", "X"), ("Y", "Y"), ("KIND", "KIND"), ("IDX", "IDX"), ("POL", "POL"), ("SUB", "SUB"),
    ("NET_A", "NET_A"), ("NET_B", "NET_B"),
    ("RULE_ID", "RULE_ID"), ("ID_PATTERN(그룹)", "_ID_PATTERN"), ("ITEM_CNT(그룹)", "_ITEM_CNT"),
    ("LSL", "LSL"), ("USL", "USL"), ("UNIT", "UNIT"), ("열 번호", "COL_NO"),
]
NUMERIC = {"X", "Y", "IDX", "_ITEM_CNT", "LSL", "USL", "COL_NO"}

FONT = Font(name="Arial", size=10)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="305496")
RAW_FILL = PatternFill("solid", fgColor="F2F2F2")          # 원본 열 구분


def read_csv(path):
    """UTF-8 / CP949(Excel 저장), 쉼표 / 탭 구분을 모두 읽는다."""
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    delim = "\t" if "\t" in text.splitlines()[0] else ","
    return list(csv.DictReader(text.splitlines(), delimiter=delim))


def to_num(v):
    """숫자 문자열은 숫자 칸으로 (정수는 int). 빈칸 / 문자는 그대로."""
    try:
        x = float(v)
    except (TypeError, ValueError):
        return v
    return int(x) if x.is_integer() and "." not in str(v) else x


def cell_value(key, v):
    return to_num(v) if key in NUMERIC and v not in ("", None) else v


def write_sheet(ws, headers, rows, widths, raw_first=False):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
        if raw_first:
            row[0].fill = RAW_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def main():
    items = read_csv(MAP_CSV)
    rules = {r["RULE_ID"]: r for r in read_csv(RULE_CSV)}
    for it in items:
        r = rules.get(it["RULE_ID"], {})
        it["_ID_PATTERN"], it["_ITEM_CNT"] = r.get("ID_PATTERN", ""), r.get("ITEM_CNT", "")

    wb = Workbook()
    ws = wb.active
    ws.title = "헤더분해"
    rows = [[cell_value(k, it.get(k, "")) for _, k in SPLIT_COLS] for it in items]
    widths = [max(len(t), 8) + 2 for t, _ in SPLIT_COLS]
    widths[0] = max(len(it["RAW_HEADER"]) for it in items) + 2
    widths[[k for _, k in SPLIT_COLS].index("ITEM_ID")] = 24
    widths[[k for _, k in SPLIT_COLS].index("_ID_PATTERN")] = 36
    write_sheet(ws, [t for t, _ in SPLIT_COLS], rows, widths, raw_first=True)

    rws = wb.create_sheet("SpecRule")
    rule_rows = read_csv(RULE_CSV)
    fields = list(rule_rows[0].keys())
    rnum = {"ITEM_CNT", "MEAN", "STDEV", "MAX", "MIN", "LSL", "USL", "PRIORITY", "N", "OUTLIER_CNT"}
    rws_rows = [[to_num(r[f]) if f in rnum else r[f] for f in fields] for r in rule_rows]
    write_sheet(rws, fields, rws_rows, [max(len(f), 8) + 2 if f not in ("ID_PATTERN", "LIMIT_BASIS", "NOTE")
                                       else 40 for f in fields])

    wb.save(OUT_XLSX)
    print(f"{OUT_XLSX}: 헤더분해 {len(items)}행 / SpecRule {len(rule_rows)}행")


if __name__ == "__main__":
    sys.exit(main())
