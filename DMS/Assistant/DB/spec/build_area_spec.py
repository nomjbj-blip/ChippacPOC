# -*- coding: utf-8 -*-
"""
Common / DMS / YMS 테이블 명세서(xlsx) 생성 (2026-10-15).
입력: extract_table_meta.sql 결과(tbl / col / cons / idx / seq / et_item.csv, def.txt) + 영역별 설명 JSON(desc_<영역>.json).
영역: 테이블 이름 접두어 기준 - Common = TQC_ / DMS = TQD_ / YMS = TQP_(TQP_GDS_ 제외) + TD_.
     소스에서 쓰지 않는 테이블(EXCLUDE)은 넣지 않는다.
설명 JSON 형식: {"tables": {테이블: {"desc", "basis", "used_by"}}, "columns": {테이블: {컬럼 또는 접두어*: [설명, 근거]}},
               "relations": [{"child", "child_cols", "parent", "parent_cols", "note"}], "notes": [...]}
반복 컬럼: 같은 접두어 + 번호가 5개 이상 이어지고 JSON 에 '접두어*' 키가 있으면 한 줄로 묶는다(예: BIN001 ~ BIN151).
TD_*: JSON 에 없는 컬럼은 Parameter 컬럼. TD_<Device>_<공정> 이 E-Test 등록 Parameter(TQP_GDS_ET_ITEM)와 맞으면 컬럼별로,
     아니면 'PARAM*' 설명으로 한 줄로 묶는다.
사용: python build_area_spec.py <CSV 폴더> <common|dms|yms> <설명 JSON> <출력 xlsx>  (보통 Make-TableSpec.ps1 이 부른다)
"""
import datetime
import glob
import json
import os
import re
import sys
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import spec_lib as L

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '..'))
QUERY_DIR = os.path.join(ROOT, r'DMS\08_DBQuery\ORACLE')
TODAY = datetime.date.today().strftime('%Y-%m-%d')

AREAS = {
    'common': {
        'label': 'Common', 'prefixes': ('TQC_',),
        'scope': 'DMS / YMS 가 함께 쓰는 공통 테이블 (사용자 / 권한 / 화면 기능 / 설정 / 설비 / 공정 / Lot 상태 / 서비스 로그)',
        'xml_prefixes': ('TQC_',),
    },
    'dms': {
        'label': 'DMS', 'prefixes': ('TQD_',),
        'scope': 'Defect 관리 (검사 Step / Wafer / Setup / Defect / Image / 검사 요약)',
        'xml_prefixes': ('TQD_',),
    },
    'yms': {
        'label': 'YMS', 'prefixes': ('TQP_', 'TD_'),
        'scope': '수율 관리 (Lot / Wafer / Bin / Program / Map / Spec) + Program 별 Die 측정 데이터(TD_*)',
        'xml_prefixes': ('TQP_', 'T_AVI_', 'T_PRB'),
    },
}
# 소스(Query XML / C#)에서 쓰지 않는 테이블 - 명세서에서 뺀다 (2026-10-15 조사)
#   TQC_SEC_GROUP: 같은 이름의 DSL 클래스는 있지만 TQC_SECURITY_GROUP.xml 을 쓴다(이 테이블을 읽는 쿼리 없음)
EXCLUDE = {'G_FAB2YMSWEH_TBL', 'G_PCM_MAST', 'PARAPROG_SPEC', 'TOAD_PLAN_TABLE', 'TQD_DEFECT_DIE_SUM', 'TQP_ZONE', 'TQC_SEC_GROUP'}
BASIS_ORDER = ['DB 주석', '원본 정의서', '설계 문서', '소스', 'E-Test 등록', '이름 추정']
BASIS_TEXT = {
    'DB 주석': 'DB 컬럼 주석',
    '원본 정의서': '같은 제품을 쓴 다른 프로젝트의 테이블 정의서 (컬럼이 같은 것만)',
    '설계 문서': '이 프로젝트 설계 문서 (문서/설계)',
    '소스': 'Query XML / C# 소스에서 쓰임을 확인',
    'E-Test 등록': 'E-Test 등록 Parameter(TQP_GDS_ET_ITEM)의 헤더 원문',
    '이름 추정': '근거 자료 없이 컬럼 이름과 값으로 추정 - 확인 필요',
}


def in_area(table, area):
    if table in EXCLUDE or table.startswith('TQP_GDS'):
        return False
    return table.startswith(AREAS[area]['prefixes'])


def query_xml(table):
    return table + '.xml' if os.path.exists(os.path.join(QUERY_DIR, table + '.xml')) else ''


def sequence_usage():
    """시퀀스 -> 그 시퀀스의 NEXTVAL 을 쓰는 Query XML 이름 목록."""
    use = defaultdict(list)
    for path in sorted(glob.glob(os.path.join(QUERY_DIR, '*.xml'))):
        text = open(path, 'rb').read().decode('utf-8-sig', 'replace')
        for m in set(re.findall(r'\b([A-Z_0-9]+)\.NEXTVAL', text, re.I)):
            use[m.upper()].append(os.path.splitext(os.path.basename(path))[0])
    return use


def column_rows(table, cols, desc_cols, et_items):
    """
    컬럼을 명세서 행으로 바꾼다. 반환: [(번호 표시, 컬럼명 표시, [원본 컬럼 행...], 설명, 근거, 묶음 여부)]
    """
    names = [c[2] for c in cols]
    result = []
    i = 0
    is_td = table.startswith('TD_')
    params = []
    while i < len(cols):
        c = cols[i]
        name = c[2]
        if is_td and name not in desc_cols:
            params.append(c)
            i += 1
            continue
        m = re.match(r'^(.*?[A-Z_])(\d+)$', name)
        if m and (m.group(1) + '*') in desc_cols:
            # 같은 접두어 + 번호가 1씩 이어지는 동안 묶는다 (USER_CMF_9 -> USER_CMF_10 처럼 자릿수가 바뀌어도 이어짐)
            prefix, prev = m.group(1), int(m.group(2)) - 1
            j = i
            while j < len(cols):
                mm = re.match(r'^(.*?[A-Z_])(\d+)$', cols[j][2])
                if not mm or mm.group(1) != prefix or int(mm.group(2)) != prev + 1:
                    break
                prev = int(mm.group(2))
                j += 1
            if j - i >= 5:
                d, b = desc_cols[prefix + '*']
                run = cols[i:j]
                result.append(('%s ~ %s' % (run[0][1], run[-1][1]), '%s ~ %s (%d개)' % (run[0][2], run[-1][2], len(run)), run, d, b, True))
                i = j
                continue
        d, b = desc_cols.get(name, ('', ''))
        result.append((int(c[1]), name, [c], d, b, False))
        i += 1
    if params:
        items = et_items.get(table, {})
        if items and all(p[2] in items for p in params):
            for p in params:
                header, unit = items[p[2]]
                text = 'E-Test 측정값: ' + header + (' (단위 %s)' % unit if unit else '')
                result.append((int(p[1]), p[2], [p], text, 'E-Test 등록', False))
        else:
            d, b = desc_cols.get('PARAM*', ('Parameter 측정값 (컬럼 이름 = Parameter 이름)', '이름 추정'))
            result.append(('%s ~ %s' % (params[0][1], params[-1][1]), 'Parameter 컬럼 (%d개)' % len(params), params, d, b, True))
        result.sort(key=lambda r: int(str(r[0]).split(' ')[0]))
    return result


def build(csv_dir, area, desc_path, out):
    conf = AREAS[area]
    desc = json.load(open(desc_path, encoding='utf-8'))
    desc_tables = desc.get('tables', {})
    desc_cols_all = desc.get('columns', {})
    relations = desc.get('relations', [])

    tbl_rows = L.load_csv(csv_dir, 'tbl.csv')
    tables = [r[0] for r in tbl_rows if in_area(r[0], area)]
    tables.sort()
    db_comment = {r[0]: (r[1] if len(r) > 1 else '') for r in tbl_rows}
    columns = defaultdict(list)
    for r in L.load_csv(csv_dir, 'col.csv'):
        columns[r[0]].append(r)
    cons = defaultdict(list)
    for r in L.load_csv(csv_dir, 'cons.csv'):
        cons[r[0]].append(r)
    indexes = defaultdict(list)
    for r in L.load_csv(csv_dir, 'idx.csv'):
        if r[3] != 'LOB':
            indexes[r[0]].append(r)
    defaults = L.load_defaults(csv_dir)
    seqs = L.load_csv(csv_dir, 'seq.csv')
    et_items = defaultdict(dict)
    for r in L.load_csv(csv_dir, 'et_item.csv'):
        et_items['TD_%s_%s' % (r[0].upper(), r[1].upper())][r[2]] = (r[3] if len(r) > 3 else '', r[4] if len(r) > 4 else '')

    warnings = []
    pk_cols = {}
    for t in tables:
        pk = [r for r in cons[t] if r[2] == 'P']
        pk_cols[t] = [x.strip() for x in pk[0][9].split(',')] if pk else []
    parents_of = defaultdict(list)   # (자식 테이블, 컬럼) -> 부모 표시
    related = defaultdict(set)
    for rel in relations:
        child, parent = rel.get('child', ''), rel.get('parent', '')
        for c in [x.strip() for x in rel.get('child_cols', '').split(',') if x.strip()]:
            label = parent if parent else ''
            if label and label not in parents_of[(child, c)]:
                parents_of[(child, c)].append(label)
        related[child].add(parent)
        related[parent].add(child)

    wb = Workbook()
    sheet_info = OrderedDict()

    # ---------------------------------------------------------------- 테이블별 시트
    for t in tables:
        ws = wb.create_sheet(t[:31])
        L.setup_print(ws)
        for i, w in enumerate([8, 28, 58, 20, 6, 5, 24, 14, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.column_dimensions['J'].hidden = True
        L.title(ws, 'A1:I1', '테이블 명세서')
        back = ws['I2']
        back.value = '<- 테이블 목록'
        L.link(back, '테이블 목록')
        back.alignment = Alignment(horizontal='right')

        td = desc_tables.get(t, {})
        tdesc = td.get('desc') or (db_comment.get(t) if db_comment.get(t) and '?' not in db_comment.get(t) else '')
        if not tdesc:
            warnings.append('테이블 설명 없음: ' + t)
        pk = [r for r in cons[t] if r[2] == 'P']
        pk_text = '%s (%s)' % (pk[0][1], pk[0][9]) if pk else '(없음)'

        L.merge(ws, 'A3:B3', '테이블명', bold=True, fill=L.LABEL_FILL)
        L.cell(ws, 'C3', t, bold=True)
        L.cell(ws, 'D3', '영역', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'E3:F3', conf['label'])
        L.cell(ws, 'G3', '스키마', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'H3:I3', 'DMSMGR')
        L.merge(ws, 'A4:B4', '테이블 설명', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'C4:I4', tdesc)
        L.fit(ws, 4, (tdesc, 159))
        L.merge(ws, 'A5:B5', '기본 키 (PK)', bold=True, fill=L.LABEL_FILL)
        L.cell(ws, 'C5', pk_text)
        L.cell(ws, 'D5', '컬럼 수', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'E5:F5', None, align=L.CENTER)
        L.cell(ws, 'G5', 'Query XML', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'H5:I5', query_xml(t))
        used_by = td.get('used_by', '')
        L.merge(ws, 'A6:B6', '사용 기능', bold=True, fill=L.LABEL_FILL)
        L.merge(ws, 'C6:I6', used_by)
        L.fit(ws, 6, (used_by, 159))

        L.header_row(ws, 8, ['No', '컬럼명', '컬럼 설명', '데이터 타입', 'NULL', 'PK', '참조 (논리)', '기본값', '근거'])
        row = 9
        desc_cols = {k: tuple(v) for k, v in desc_cols_all.get(t, {}).items()}
        for no, name, members, d, b, grouped in column_rows(t, columns[t], desc_cols, et_items):
            first = members[0]
            if not d:
                comment = first[10] if len(first) > 10 else ''
                if comment and '?' not in comment:
                    d, b = comment, 'DB 주석'
            if not d:
                warnings.append('컬럼 설명 없음: %s.%s' % (t, name))
            types = sorted(set(L.data_type(m[3], m[5], m[6], m[7], m[8]) for m in members))
            nulls = sorted(set('N' if m[9] == 'N' else 'Y' for m in members))
            dfs = sorted(set(defaults.get((t, m[2]), '') for m in members))
            pk_pos = '' if grouped else (pk_cols[t].index(name) + 1 if name in pk_cols[t] else '')
            refs = '' if grouped else ', '.join(parents_of.get((t, name), []))
            fill = L.GROUP_FILL if grouped else (L.PK_FILL if pk_pos else None)
            L.write_row(ws, row, [no, name, d, ' / '.join(types), '/'.join(nulls), pk_pos, refs,
                                  ' / '.join(x for x in dfs if x), b, len(members)], fill=fill, centers=(1, 5, 6, 9))
            row += 1
        first_row, last_row = 9, row - 1
        ws['E5'].value = '=SUM(J%d:J%d)' % (first_row, last_row)

        row += 1
        L.merge(ws, 'A%d:I%d' % (row, row), '제약 조건', bold=True, fill=L.SECTION_FILL)
        row += 1
        L.header_row(ws, row, ['No', '구분', '제약 이름', '컬럼', '', '', '내용', '', ''])
        ws.merge_cells('D%d:F%d' % (row, row))
        ws.merge_cells('G%d:I%d' % (row, row))
        row += 1
        kinds = {'P': ('PK', '기본 키'), 'U': ('UNIQUE', '중복 금지'), 'R': ('FK', ''), 'C': ('CHECK', '')}
        user_cons = [r for r in cons[t] if not (r[2] == 'C' and r[8] == 'GENERATED NAME')]
        if not user_cons:
            L.write_row(ws, row, ['-', '-', '(없음)', '', '', '', '', '', ''], centers=(1, 2))
            ws.merge_cells('D%d:F%d' % (row, row))
            ws.merge_cells('G%d:I%d' % (row, row))
            row += 1
        for n, r in enumerate(user_cons, 1):
            kind, content = kinds.get(r[2], (r[2], ''))
            if r[2] == 'C':
                content = L.clean(r[3])
            elif r[2] == 'R':
                content = '-> %s (삭제 규칙 %s)' % (r[5], r[6])
            L.write_row(ws, row, [n, kind, r[1], r[9] if r[2] != 'C' else '', '', '', content, '', ''], centers=(1, 2))
            ws.merge_cells('D%d:F%d' % (row, row))
            ws.merge_cells('G%d:I%d' % (row, row))
            L.fit(ws, row, (content, 49), (r[9], 31))
            row += 1

        row += 1
        L.merge(ws, 'A%d:I%d' % (row, row), '인덱스', bold=True, fill=L.SECTION_FILL)
        row += 1
        L.header_row(ws, row, ['No', '유일', '인덱스 이름', '컬럼', '', '', '용도', '', ''])
        ws.merge_cells('D%d:F%d' % (row, row))
        ws.merge_cells('G%d:I%d' % (row, row))
        row += 1
        if not indexes[t]:
            L.write_row(ws, row, ['-', '-', '(없음)', '', '', '', '', '', ''], centers=(1, 2))
            ws.merge_cells('D%d:F%d' % (row, row))
            ws.merge_cells('G%d:I%d' % (row, row))
            row += 1
        for n, r in enumerate(indexes[t], 1):
            use = '기본 키' if 'P' in r[5] else ('중복 금지' if 'U' in r[5] else '')
            L.write_row(ws, row, [n, 'UNIQUE' if r[2] == 'UNIQUE' else '-', r[1], r[4], '', '', use, '', ''], centers=(1, 2))
            ws.merge_cells('D%d:F%d' % (row, row))
            ws.merge_cells('G%d:I%d' % (row, row))
            L.fit(ws, row, (r[4], 31))
            row += 1
        ws.freeze_panes = 'A9'
        sheet_info[t] = (ws.title, first_row, last_row, tdesc, pk[0][9] if pk else '', used_by)

    present_basis = [b for b in BASIS_ORDER if any(
        wb[s[0]].cell(r, 9).value == b for s in sheet_info.values() for r in range(s[1], s[2] + 1))]

    # ---------------------------------------------------------------- 테이블 목록
    ws = wb.create_sheet('테이블 목록', 0)
    L.setup_print(ws)
    widths = [5, 24, 50, 30, 8, 36, 30, 24] + [9] * len(present_basis)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last_col = get_column_letter(len(widths))
    L.title(ws, 'A1:%s1' % last_col, '%s 테이블 목록' % conf['label'])
    L.header_row(ws, 3, ['No', '테이블명', '테이블 설명', '기본 키 (PK)', '컬럼 수', '사용 기능', '관련 테이블 (논리)', 'Query XML'] + ['근거: ' + b for b in present_basis])
    row = 4
    for n, (t, (sname, fr, lr, tdesc, pkc, used_by)) in enumerate(sheet_info.items(), 1):
        rel = ', '.join(sorted(x for x in related.get(t, set()) if x and x != t))
        values = [n, t, tdesc, pkc, "='%s'!E5" % sname, used_by, rel, query_xml(t)]
        values += ['=COUNTIF(\'%s\'!I%d:I%d,"%s")' % (sname, fr, lr, b) for b in present_basis]
        L.write_row(ws, row, values, centers=tuple([1, 5] + list(range(9, 9 + len(present_basis)))))
        L.link(ws.cell(row=row, column=2), sname)
        row += 1
    list_first, list_last = 4, row - 1
    L.cell(ws, 'D%d' % row, '합계', bold=True, fill=L.LABEL_FILL, align=L.CENTER)
    L.cell(ws, 'E%d' % row, '=SUM(E%d:E%d)' % (list_first, list_last), bold=True, fill=L.LABEL_FILL, align=L.CENTER)
    for k in range(len(present_basis)):
        col = get_column_letter(9 + k)
        L.cell(ws, '%s%d' % (col, row), '=SUM(%s%d:%s%d)' % (col, list_first, col, list_last), bold=True, fill=L.LABEL_FILL, align=L.CENTER)
    total_row = row
    ws.freeze_panes = 'C4'

    # ---------------------------------------------------------------- 관계 (논리)
    ws = wb.create_sheet('관계(논리)')
    L.setup_print(ws)
    for i, w in enumerate([5, 24, 30, 24, 30, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    L.title(ws, 'A1:F1', '테이블 관계 (논리)')
    L.merge(ws, 'A2:F2', '이 영역의 테이블에는 DB FK 가 없다. Query XML / 소스의 조인 조건으로 확인한 관계만 적는다. [다른 영역] = 부모 또는 자식이 이 명세서 밖의 테이블.', border=False)
    L.header_row(ws, 4, ['No', '자식 테이블', '자식 컬럼', '부모 테이블', '부모 컬럼', '근거 / 비고'])
    row = 5
    n = 0
    for rel in relations:
        child, parent = rel.get('child', ''), rel.get('parent', '')
        if not (in_area(child, area) or in_area(parent, area)):
            continue
        n += 1
        mark = '' if (in_area(child, area) and in_area(parent, area)) else '[다른 영역] '
        L.write_row(ws, row, [n, child, rel.get('child_cols', ''), parent, rel.get('parent_cols', ''), mark + rel.get('note', '')], centers=(1,))
        row += 1
    if n == 0:
        L.write_row(ws, row, ['-', '(확인된 관계 없음)', '', '', '', ''])
    ws.freeze_panes = 'A5'

    # ---------------------------------------------------------------- 시퀀스
    ws = wb.create_sheet('시퀀스')
    L.setup_print(ws)
    for i, w in enumerate([5, 20, 46, 8, 8, 8, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    L.title(ws, 'A1:G1', '시퀀스')
    L.header_row(ws, 3, ['No', '시퀀스', '사용 Query XML', '시작', '증가', 'CACHE', '비고'])
    use = sequence_usage()
    row = 4
    n = 0
    area_of_xml = lambda x: next((a for a, c in AREAS.items() if x.startswith(c['xml_prefixes']) and not x.startswith('TQP_GDS')), '')
    for s in seqs:
        xmls = [x for x in use.get(s[0], []) if not x.startswith('TQP_GDS')]
        areas = sorted(set(area_of_xml(x) for x in xmls) - {''})
        if area not in areas:
            continue
        n += 1
        note = ''
        if len(areas) > 1:
            note = ' / '.join(AREAS[a]['label'] for a in areas) + ' 가 같은 시퀀스를 쓴다 (번호가 서로 겹치지 않음)'
        L.write_row(ws, row, [n, s[0], ', '.join(x + '.xml' for x in xmls), int(s[1]), int(s[2]), int(s[3]), note], centers=(1, 4, 5, 6))
        row += 1
    if n == 0:
        L.write_row(ws, row, ['-', '(없음)', '', '', '', '', ''])
    # POC 시퀀스(TQP_GDS_* 용)는 POC 테이블 명세서에서 다룬다
    poc_seqs = {'IMPORT_JOB_SEQ', 'CHAIN_SEQ', 'BATCH_EDIT_SEQ', 'SRC_FILE_SEQ', 'ET_ITEM_SEQ'}
    unused = [s[0] for s in seqs if not use.get(s[0]) and s[0] not in poc_seqs]

    # ---------------------------------------------------------------- 개요
    ws = wb['Sheet']
    ws.title = '개요'
    wb.move_sheet(ws, offset=-wb.index(ws))
    L.setup_print(ws, landscape=False)
    for i, w in enumerate([4, 22, 90], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    L.merge(ws, 'B2:C2', 'SCK 분석시스템 구축 POC - %s 테이블 명세서' % conf['label'], bold=True, size=16, align=Alignment(vertical='center'), border=False)
    ws.row_dimensions[2].height = 32
    excluded = sorted(t for t in EXCLUDE if t.startswith(conf['prefixes']))
    info = [
        ('대상 시스템', 'DACrux DMS / YMS (SCK 분석시스템 구축 POC)'),
        ('영역', '%s - %s' % (conf['label'], conf['scope'])),
        ('기준 DB', '로컬 개발 DB (Oracle 19c, 스키마 DMSMGR, 문자셋 AL32UTF8)'),
        ('조회 일자', TODAY),
        ('작성 근거', 'DB 메타데이터 (컬럼 / 주석 / 제약 / 인덱스 / 시퀀스) + Query XML (DMS/08_DBQuery/ORACLE) / C# 소스 + 같은 제품을 쓴 다른 프로젝트의 테이블 정의서 + 이 프로젝트 설계 문서'),
        ('대상 범위', '테이블 %d개 (이름 접두어 %s). POC 신규 테이블(TQP_GDS_*)은 POC 테이블 명세서에 따로 있다.' % (len(tables), ' / '.join(conf['prefixes']))),
    ]
    if excluded:
        info.append(('제외', '소스(Query XML / C#)에서 쓰지 않는 테이블: ' + ', '.join(excluded)))
    row = 4
    for k, v in info:
        L.cell(ws, 'B%d' % row, k, bold=True, fill=L.LABEL_FILL)
        L.cell(ws, 'C%d' % row, v)
        row += 1
    row += 1
    L.cell(ws, 'B%d' % row, '집계', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    L.cell(ws, 'C%d' % row, '값', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    row += 1
    c = L.cell(ws, 'B%d' % row, '테이블 수', bold=True)
    c = L.cell(ws, 'C%d' % row, "=COUNTA('테이블 목록'!B%d:B%d)" % (list_first, list_last), align=Alignment(horizontal='left', vertical='center'))
    c.number_format = '0"개"'
    row += 1
    L.cell(ws, 'B%d' % row, '컬럼 수', bold=True)
    c = L.cell(ws, 'C%d' % row, "='테이블 목록'!E%d" % total_row, align=Alignment(horizontal='left', vertical='center'))
    c.number_format = '0"개"'
    row += 1
    for k, b in enumerate(present_basis):
        L.cell(ws, 'B%d' % row, '근거: ' + b, bold=True)
        c = L.cell(ws, 'C%d' % row, "='테이블 목록'!%s%d" % (get_column_letter(9 + k), total_row), align=Alignment(horizontal='left', vertical='center'))
        c.number_format = '0"행  - %s"' % BASIS_TEXT[b]
        row += 1
    row += 1
    L.cell(ws, 'B%d' % row, '시트', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    L.cell(ws, 'C%d' % row, '내용', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    row += 1
    for sheet, v in [('테이블 목록', '테이블 요약 (테이블명을 누르면 해당 시트로 이동, 근거별 행 수)'),
                     ('테이블별 시트', '컬럼 / 제약 조건 / 인덱스'),
                     ('관계(논리)', 'Query XML / 소스의 조인으로 확인한 테이블 관계 (DB FK 없음)'),
                     ('시퀀스', '이 영역 테이블이 쓰는 시퀀스')]:
        c = L.cell(ws, 'B%d' % row, sheet, bold=True)
        if sheet in wb.sheetnames:
            L.link(c, sheet)
            c.font = L.font(True, 10, '0563C1', 'single')
        L.cell(ws, 'C%d' % row, v)
        row += 1
    row += 1
    L.cell(ws, 'B%d' % row, '표기 기준', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    L.cell(ws, 'C%d' % row, '설명', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
    row += 1
    legend = [('데이터 타입', 'VARCHAR2(n CHAR) = 글자 수 기준, VARCHAR2(n) = byte 기준'),
              ('NULL / PK', 'NULL: N = NOT NULL (필수), Y = NULL 허용. PK: 기본 키 안 순서 (노란 행)'),
              ('반복 컬럼', '같은 접두어 + 번호가 이어지는 컬럼은 한 줄로 묶었다 (회색 행, 예: BIN001 ~ BIN151). No 는 첫 ~ 끝 컬럼 순번'),
              ('참조 (논리)', '조인으로 확인한 부모 테이블. DB FK 는 없다'),
              ('근거', ' / '.join('%s = %s' % (b, BASIS_TEXT[b]) for b in present_basis))]
    for k, v in legend:
        L.cell(ws, 'B%d' % row, k, bold=True)
        L.cell(ws, 'C%d' % row, v)
        row += 1
    notes = list(desc.get('report_notes', []))
    # 업로드 때 생기는 TD_* 데이터 테이블은 조회 시점에 없을 수 있다 (설명은 JSON 에 남겨 다시 생기면 자동으로 들어감)
    absent_td = sorted(t for t in desc_tables if t.startswith('TD_') and t not in tables)
    if absent_td:
        notes.append('조회 시점에 DB 에 없는 데이터 테이블: %s (E-Test 업로드 때 생기고 데이터를 지우면 함께 삭제되는 TD_* 테이블)' % ', '.join(absent_td))
    no_pk = [t for t in tables if not pk_cols[t]]
    if no_pk:
        notes.append('DB 에 기본 키(PK)가 없는 테이블 %d개: %s' % (len(no_pk), ', '.join(no_pk)))
    if area == 'common' and unused:
        notes.append('쓰지 않는 시퀀스 %d개(%s)는 시퀀스 시트에서 뺐다.' % (len(unused), ', '.join(unused)))
    if notes:
        row += 1
        L.cell(ws, 'B%d' % row, '참고', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
        L.cell(ws, 'C%d' % row, '내용', bold=True, fill=L.HEAD_FILL, color='FFFFFF', align=L.CENTER)
        row += 1
        for n, v in enumerate(notes, 1):
            L.cell(ws, 'B%d' % row, '참고 %d' % n, bold=True)
            L.cell(ws, 'C%d' % row, v)
            row += 1

    wb.active = 0
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    total_cols = sum(len(columns[t]) for t in tables)
    print('저장:', out)
    print('%s: 테이블 %d개 / 컬럼 %d개 / 근거 %s' % (conf['label'], len(tables), total_cols, ', '.join(present_basis)))
    unknown = sorted(t for t in set(desc_tables) - set(tables) if not t.startswith('TD_'))
    if unknown:
        warnings.append('설명 JSON 에만 있는 테이블: ' + ', '.join(unknown))
    for w in warnings:
        print('경고:', w)


if __name__ == '__main__':
    if len(sys.argv) < 5:
        sys.exit('사용: python build_area_spec.py <CSV 폴더> <common|dms|yms> <설명 JSON> <출력 xlsx>')
    build(sys.argv[1], sys.argv[2].lower(), sys.argv[3], sys.argv[4])
