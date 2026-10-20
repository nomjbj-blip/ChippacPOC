# -*- coding: utf-8 -*-
"""
POC 추가 테이블 명세서(xlsx) 생성 (2026-10-15).
입력: extract_table_meta.sql 로 뽑은 DB 메타데이터(tbl / col / cons / idx.csv, def.txt) + Persistence DDL(*.sql).
설명: 컬럼 설명은 DESC 사전 -> DB 주석 -> SQL COMMENT ON 원문 -> 공통 컬럼 설명 순서로 쓴다.
     DB 주석이 없는 GDS Map / Chain 컬럼은 DESC 사전(설계 문서 기준)에 적는다.
손으로 맞출 목록: TABLES(대상 테이블 / 업무 구분 / 설명), DESC, APPLY_ORDER(같은 날 SQL 적용 순서),
     SEQS / DROPPED / EXISTING / HISTORY(변경 이력 시트). 새 테이블이나 DDL 이 생기면 여기에 더한다.
사용: python build_table_spec.py <CSV 폴더> <출력 xlsx>  (보통 Make-TableSpec.ps1 이 부른다)
"""
import csv, datetime, glob, os, re, sys
from collections import defaultdict, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 이 파일 위치: <ChippacPOC>\DMS\Assistant\DB\spec
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '..'))
PERSIST = os.path.join(ROOT, r'DMS\06_CoreLibSource\GdsMapControl\Persistence')
if len(sys.argv) < 3:
    sys.exit('사용: python build_table_spec.py <CSV 폴더> <출력 xlsx>')
CSV_DIR, OUT = sys.argv[1], sys.argv[2]
TODAY = datetime.date.today().strftime('%Y-%m-%d')

def load(name):
    # 추출 파일은 DMSMGR 전체 테이블이므로 POC(TQP_GDS_*) 테이블만 쓴다
    return [r for r in csv.reader(open(os.path.join(CSV_DIR, name), encoding='utf-8'), delimiter='|')
            if r and r[0] != 'TABLE_NAME' and r[0].startswith('TQP_GDS')]

cols_raw, cons_raw, idx_raw, tbl_raw = load('col.csv'), load('cons.csv'), load('idx.csv'), load('tbl.csv')
defaults = {}
for line in open(os.path.join(CSV_DIR, 'def.txt'), encoding='utf-8'):
    p = line.rstrip('\n').split('|', 2)
    if len(p) == 3 and p[0].startswith('TQP_GDS'):
        defaults[(p[0], p[1])] = p[2].strip()
VIRTUAL = {('TQP_GDS_SRC_FILE', 'WAFER_ID')}

# ---------------------------------------------------------------- 테이블 순서 / 구분 / 설명
TABLES = OrderedDict([
    ('TQP_GDS_MAP', ('GDS Map', 'GDS Map 헤더. Factory + Device 당 1건으로 GDS 파일 / Top Structure / 단위 / 저장 상태 / 좌표 기준면(Front / Back)을 관리한다.')),
    ('TQP_GDS_LAYER', ('GDS Map', 'GDS Map 의 Layer 정의 (Layer 번호 / 기본 색 / 표시 이름).')),
    ('TQP_GDS_PLACED_EL', ('GDS Map', '반복 배치(SREF)를 펼친 실제 도형. 월드 좌표 / Bounds / 좌표 배열 / TEXT 를 저장하며 Map 화면 / Chain / Label 검색 / 연관 분석이 함께 쓴다.')),
    ('TQP_GDS_IMPORT_JOB', ('GDS Map', 'GDS Map 대량 저장 작업. Device 당 1건으로 예상 / 수신 건수와 진행 상태를 관리한다.')),
    ('TQP_GDS_IMPORT_BATCH', ('GDS Map', 'Map 저장 작업의 배치 전송 기록 (Layer / 배치 도형). 중복 전송 판정과 완료 확인에 쓴다.')),
    ('TQP_GDS_CHAIN', ('GDS Chain', 'Chain 마스터. Chain 번호 / Code / 사용 상태를 관리한다.')),
    ('TQP_GDS_CHAIN_VER', ('GDS Chain', 'Chain 설정 (이름 / 색 / 탐색 옵션)과 저장 상태. 현재는 버전 이력 없이 VERSION_NO = 1 로 덮어쓴다.')),
    ('TQP_GDS_CHAIN_MEMBER', ('GDS Chain', 'Chain 에 포함된 도형 (Input / Output / Member)과 단자 클릭 위치.')),
    ('TQP_GDS_CHAIN_LAYER', ('GDS Chain', 'Chain 탐색 때 선택한 Layer.')),
    ('TQP_GDS_CHAIN_RULE', ('GDS Chain', 'Chain 탐색의 Layer 조합 연결 규칙과 허용 거리.')),
    ('TQP_GDS_CHAIN_EDIT', ('GDS Chain', '도형 수동 추가 / 제외 / 복원 이력. 다시 탐색할 때 엔지니어가 정한 내용을 다시 적용한다.')),
    ('TQP_GDS_SRC_FILE', ('E-Test', 'E-Test 결과 파일 1개 = 1행. 파일명 해석 값 / 적재 건수 / Spec 판정 결과 / Map 좌표 계산 / YMS 반영 상태를 관리한다.')),
    ('TQP_GDS_ET_ITEM', ('E-Test', 'E-Test 측정 항목(Parameter) 마스터. Device + 공정 단위 등록 Parameter 와 GDS Chain 이름 연결을 관리한다.')),
    ('TQP_GDS_ET_DIE', ('E-Test', 'E-Test Die 정보 (파일 행 단위). 원본 좌표 / Bin / DUT Flag / Map 좌표 / Spec 판정.')),
    ('TQP_GDS_ET_VALUE', ('E-Test', 'E-Test 측정값 (Die x 측정 항목).')),
    ('TQP_GDS_ET_LIMIT', ('E-Test', 'E-Test Spec (하한 / 상한 / 사용 여부). Device + 공정 + 등록 Parameter 단위.')),
])

# ---------------------------------------------------------------- 컬럼 설명 (DB 주석이 없거나 깨진 컬럼)
COMMON = {
    'FACTORY': 'Factory (로그인 Factory, 예: SCK)',
    'DEVICE_ID': 'Device(제품) ID. TQP_GDS_MAP 의 Device',
    'CREATE_USER': '생성 사용자 ID',
    'CREATE_TIME': '생성 일시',
    'UPDATE_USER': '수정 사용자 ID',
    'UPDATE_TIME': '수정 일시 (Service 가 저장할 때 기록)',
}
ET_DEVICE = 'Device(제품) ID. E-Test 파일 Device 와 같은 이름의 GDS Map Device'
CHAIN_CHILD = {
    'CHAIN_SEQ': 'Chain 번호 (TQP_GDS_CHAIN)',
    'VERSION_NO': 'Chain 버전 번호 (TQP_GDS_CHAIN_VER, 현재 1 고정)',
}
DESC = {
    ('TQP_GDS_MAP', 'DEVICE_ID'): 'Device(제품) ID. Factory + Device 당 Map 1건 (KLARF Product / E-Test Device 와 이름으로 연결)',
    ('TQP_GDS_MAP', 'TOP_STRUCTURE'): 'GDS Top Structure 이름 (도형을 펼친 기준 Structure)',
    ('TQP_GDS_MAP', 'GDS_FILE_NAME'): '원본 GDS 파일 이름',
    ('TQP_GDS_MAP', 'USER_UNIT'): 'GDS User Unit (HEADER UNITS 의 첫째 값)',
    ('TQP_GDS_MAP', 'DATABASE_UNIT'): 'GDS Database Unit, 미터 단위 (HEADER UNITS 의 둘째 값)',
    ('TQP_GDS_MAP', 'ID_GENERATOR_VERSION'): '도형 ID(PLACED_ELEMENT_ID) 생성 규칙 버전',
    ('TQP_GDS_MAP', 'POINTS_FORMAT_VERSION'): '좌표 배열(TQP_GDS_PLACED_EL.POINTS_BIN) 형식 버전',
    ('TQP_GDS_MAP', 'MAP_STATUS'): '저장 상태. DRAFT = 저장 중 / READY = 저장 완료 (조회 / Chain / 연관 분석은 READY 만 사용)',
    ('TQP_GDS_LAYER', 'LAYER_ID'): 'GDS Layer 번호',
    ('TQP_GDS_LAYER', 'COLOR_ARGB'): '기본 표시 색 (ARGB 정수)',
    ('TQP_GDS_LAYER', 'DISPLAY_NAME'): 'Layer 표시 이름',
    ('TQP_GDS_PLACED_EL', 'PLACED_ELEMENT_ID'): '배치 도형 ID. 펼친 배치 경로 + 원본 순번 (예: R0/E15921). Factory + Device 안에서 고유',
    ('TQP_GDS_PLACED_EL', 'SOURCE_ELEMENT_ID'): '원본 Element ID (예: TOP/E15921). 같은 원본에서 반복 배치된 도형을 묶는 값 (원본 테이블은 2026-10-12 삭제, FK 없음)',
    ('TQP_GDS_PLACED_EL', 'LAYER_ID'): 'GDS Layer 번호 (TQP_GDS_LAYER)',
    ('TQP_GDS_PLACED_EL', 'DATA_TYPE'): 'GDS DATATYPE 번호',
    ('TQP_GDS_PLACED_EL', 'ELEMENT_TYPE'): '도형 종류',
    ('TQP_GDS_PLACED_EL', 'CLOSED_FLAG'): '닫힌 도형 여부 (1 = 닫힘)',
    ('TQP_GDS_PLACED_EL', 'PATH_WIDTH'): 'PATH 폭 (PATH 도형만)',
    ('TQP_GDS_PLACED_EL', 'PATH_TYPE'): 'GDS PATHTYPE (PATH 도형만, 그 밖에는 NULL)',
    ('TQP_GDS_PLACED_EL', 'MIN_X'): '도형 외곽 최소 X (월드 좌표)',
    ('TQP_GDS_PLACED_EL', 'MIN_Y'): '도형 외곽 최소 Y (월드 좌표)',
    ('TQP_GDS_PLACED_EL', 'MAX_X'): '도형 외곽 최대 X (월드 좌표)',
    ('TQP_GDS_PLACED_EL', 'MAX_Y'): '도형 외곽 최대 Y (월드 좌표)',
    ('TQP_GDS_PLACED_EL', 'POINT_COUNT'): '좌표 점 개수',
    ('TQP_GDS_PLACED_EL', 'POINTS_FORMAT_VERSION'): '좌표 배열(POINTS_BIN) 형식 버전',
    ('TQP_GDS_PLACED_EL', 'POINTS_BIN'): '좌표 배열. 점마다 X / Y double 2개 (little-endian 16 byte)',
    ('TQP_GDS_PLACED_EL', 'TEXT_VALUE'): 'TEXT 문자열 (TEXT 도형만)',
    ('TQP_GDS_PLACED_EL', 'TEXT_TYPE'): 'GDS TEXTTYPE 번호 (TEXT 도형만)',
    ('TQP_GDS_PLACED_EL', 'TEXT_FONT'): 'GDS TEXT 글꼴 번호 (TEXT 도형만)',
    ('TQP_GDS_PLACED_EL', 'TEXT_HORIZONTAL'): 'TEXT 가로 정렬 (예: Center)',
    ('TQP_GDS_PLACED_EL', 'TEXT_VERTICAL'): 'TEXT 세로 정렬 (예: Middle)',
    ('TQP_GDS_IMPORT_JOB', 'IMPORT_JOB_SEQ'): 'Map 저장 작업 번호 (IMPORT_JOB_SEQ 시퀀스)',
    ('TQP_GDS_IMPORT_JOB', 'DEVICE_ID'): 'Device(제품) ID. Device 당 작업 1건 (UQ_GDS_IMPORT_MAP)',
    ('TQP_GDS_IMPORT_JOB', 'JOB_STATUS'): '작업 상태. UPLOADING = 전송 중 / VALIDATING = 검증 중 / READY = 완료 / FAILED = 실패 / CANCELLED = 취소',
    ('TQP_GDS_IMPORT_JOB', 'EXPECTED_PLACED_COUNT'): '보낼 예정인 도형 수',
    ('TQP_GDS_IMPORT_JOB', 'EXPECTED_POINT_COUNT'): '보낼 예정인 좌표 점 수',
    ('TQP_GDS_IMPORT_JOB', 'RECEIVED_PLACED_COUNT'): '받은 도형 수',
    ('TQP_GDS_IMPORT_JOB', 'RECEIVED_POINT_COUNT'): '받은 좌표 점 수',
    ('TQP_GDS_IMPORT_JOB', 'UPDATE_TIME'): '수정 일시 (진행 상태가 바뀔 때 기록)',
    ('TQP_GDS_IMPORT_BATCH', 'IMPORT_JOB_SEQ'): 'Map 저장 작업 번호 (TQP_GDS_IMPORT_JOB)',
    ('TQP_GDS_IMPORT_BATCH', 'STREAM_KIND'): '배치 종류. LAYER = Layer / PLACED = 배치 도형',
    ('TQP_GDS_IMPORT_BATCH', 'BATCH_NO'): '배치 번호 (종류별 0부터 빠짐없이)',
    ('TQP_GDS_IMPORT_BATCH', 'FORMAT_VERSION'): '전송 배치 형식 버전',
    ('TQP_GDS_IMPORT_BATCH', 'PAYLOAD_SHA256'): '배치 내용 SHA-256 (16진수 64자리). 같은 번호를 다시 받으면 이 값으로 같은 배치인지 비교',
    ('TQP_GDS_IMPORT_BATCH', 'ITEM_COUNT'): '배치 안 행 수',
    ('TQP_GDS_IMPORT_BATCH', 'POINT_COUNT'): '배치 안 좌표 점 수',
    ('TQP_GDS_IMPORT_BATCH', 'PAYLOAD_BYTES'): '배치 크기 (byte)',
    ('TQP_GDS_IMPORT_BATCH', 'COMMIT_TIME'): '배치 저장 확정 일시',
    ('TQP_GDS_CHAIN', 'CHAIN_SEQ'): 'Chain 번호 (CHAIN_SEQ 시퀀스). 다시 저장해도 유지',
    ('TQP_GDS_CHAIN', 'CHAIN_CODE'): 'Chain Code (A, B, ... Z, AA ...). Device 안에서 고유',
    ('TQP_GDS_CHAIN', 'CHAIN_STATUS'): 'Chain 상태. ACTIVE = 사용 / RETIRED = 사용 종료 (논리 삭제)',
    ('TQP_GDS_CHAIN', 'RETIRE_TIME'): '사용 종료 (RETIRED) 일시',
    ('TQP_GDS_CHAIN_VER', 'VERSION_NO'): 'Chain 버전 번호. 현재는 저장할 때 덮어써서 1 고정',
    ('TQP_GDS_CHAIN_VER', 'CHAIN_NAME'): 'Chain 이름 (E-Test Para 설정의 GDS Chain 이름과 연결)',
    ('TQP_GDS_CHAIN_VER', 'COLOR_ARGB'): 'Chain 표시 색 (ARGB 정수)',
    ('TQP_GDS_CHAIN_VER', 'APPLY_LAYER_RULES'): 'Layer 조합 규칙(TQP_GDS_CHAIN_RULE) 적용 여부 (1 = 적용)',
    ('TQP_GDS_CHAIN_VER', 'CELL_SIZE'): '연결 탐색 격자 크기 (화면에서 숨김, 기본 10)',
    ('TQP_GDS_CHAIN_VER', 'VERSION_STATUS'): '저장 상태. DRAFT = 저장 중 / READY = 저장 완료 / SUPERSEDED = 이전 버전',
    ('TQP_GDS_CHAIN_VER', 'IS_CURRENT'): '현재 버전 여부 (1 = 현재). READY 일 때만 1, Chain 당 최대 1개',
    ('TQP_GDS_CHAIN_MEMBER', 'PLACED_ELEMENT_ID'): 'Chain 에 포함된 도형 ID (TQP_GDS_PLACED_EL). 한 도형이 여러 Chain 에 속할 수 있음',
    ('TQP_GDS_CHAIN_MEMBER', 'ROLE_CODE'): '도형 역할. INPUT = 입력 단자 (1개 필수) / OUTPUT = 출력 단자 (0~1개) / MEMBER = 연결 도형',
    ('TQP_GDS_CHAIN_MEMBER', 'INCLUSION_SOURCE'): '포함 경위. TRACE = 연결 탐색 / MANUAL = 수동 추가 / ENDPOINT = Input / Output 단자 지정',
    ('TQP_GDS_CHAIN_MEMBER', 'MARKER_X'): '단자 클릭 위치 X (INPUT / OUTPUT 만, MEMBER 는 NULL)',
    ('TQP_GDS_CHAIN_MEMBER', 'MARKER_Y'): '단자 클릭 위치 Y (INPUT / OUTPUT 만, MEMBER 는 NULL)',
    ('TQP_GDS_CHAIN_LAYER', 'LAYER_ID'): 'Chain 탐색 때 선택한 GDS Layer 번호 (TQP_GDS_LAYER)',
    ('TQP_GDS_CHAIN_RULE', 'FIRST_LAYER_ID'): '조합 Layer 1 (작은 번호 쪽)',
    ('TQP_GDS_CHAIN_RULE', 'SECOND_LAYER_ID'): '조합 Layer 2 (큰 번호 쪽)',
    ('TQP_GDS_CHAIN_RULE', 'TOLERANCE'): '두 Layer 도형을 연결로 볼 허용 거리 (0 이상)',
    ('TQP_GDS_CHAIN_EDIT', 'EDIT_SEQ'): '편집 순번 (1부터, ADD -> EXCLUDE 순서)',
    ('TQP_GDS_CHAIN_EDIT', 'PLACED_ELEMENT_ID'): '편집 대상 도형 ID (TQP_GDS_PLACED_EL)',
    ('TQP_GDS_CHAIN_EDIT', 'ACTION_CODE'): '편집 종류. ADD = 수동 추가 / EXCLUDE = 제외 / RESTORE = 제외 취소',
    ('TQP_GDS_CHAIN_EDIT', 'BATCH_EDIT_SEQ'): '일괄 편집 묶음 번호 (BATCH_EDIT_SEQ 시퀀스, 현재 미사용 NULL)',
    ('TQP_GDS_SRC_FILE', 'SRC_FILE_SEQ'): 'E-Test 파일 번호 (SRC_FILE_SEQ 시퀀스)',
    ('TQP_GDS_SRC_FILE', 'SHEET_NAME'): '측정값을 읽은 Sheet 이름 (예: DUT Summary)',
    ('TQP_GDS_SRC_FILE', 'DIE_COUNT'): '파일의 Die 수',
    ('TQP_GDS_ET_DIE', 'SRC_FILE_SEQ'): 'E-Test 파일 번호 (TQP_GDS_SRC_FILE)',
    ('TQP_GDS_ET_DIE', 'DUT_NO'): 'DUT 번호 (파일 A열 #3 -> 3)',
    ('TQP_GDS_ET_DIE', 'DIE_X'): '파일의 Die 좌표 X (원본 그대로)',
    ('TQP_GDS_ET_DIE', 'DIE_Y'): '파일의 Die 좌표 Y (원본 그대로)',
    ('TQP_GDS_ET_DIE', 'HW_BIN'): 'Hardware Bin (예: Bin 1 -> 1)',
    ('TQP_GDS_ET_DIE', 'SW_BIN'): 'Software Bin (예: Bin 5001 -> 5001)',
    ('TQP_GDS_ET_DIE', 'TEST_TIME_MS'): 'Die 시험 시간 (ms)',
    ('TQP_GDS_ET_DIE', 'TESTS_EXECUTED'): '실행한 Test 수',
    ('TQP_GDS_ET_DIE', 'HEAD_SITE'): 'Head / Site (예: Head 1 - Site 1)',
    ('TQP_GDS_ET_ITEM', 'ET_ITEM_SEQ'): '측정 항목 번호 (ET_ITEM_SEQ 시퀀스)',
    ('TQP_GDS_ET_ITEM', 'DEVICE_ID'): ET_DEVICE,
    ('TQP_GDS_ET_ITEM', 'UNIT'): '측정 단위 (등록한 파일 기준)',
    ('TQP_GDS_ET_ITEM', 'TEST_NO'): '파일 2행 Test Number',
    ('TQP_GDS_ET_ITEM', 'DISPLAY_ORDER'): '표시 순서 (처음 등록한 파일의 열 순서)',
    ('TQP_GDS_ET_ITEM', 'REMARK'): '비고',
    ('TQP_GDS_ET_LIMIT', 'DEVICE_ID'): ET_DEVICE,
    ('TQP_GDS_ET_LIMIT', 'UNIT'): '단위',
    ('TQP_GDS_ET_LIMIT', 'REMARK'): '비고',
    ('TQP_GDS_ET_VALUE', 'SRC_FILE_SEQ'): 'E-Test 파일 번호 (TQP_GDS_ET_DIE)',
    ('TQP_GDS_ET_VALUE', 'DUT_NO'): 'DUT 번호 (TQP_GDS_ET_DIE)',
    ('TQP_GDS_ET_VALUE', 'ET_ITEM_SEQ'): '측정 항목 번호 (TQP_GDS_ET_ITEM)',
    ('TQP_GDS_ET_VALUE', 'MEAS_VALUE'): '측정값 (숫자). 값이 없으면 NULL (VALUE_STATUS 참고)',
}

# ---------------------------------------------------------------- SQL 파일: COMMENT ON 원문 / 테이블별 생성 / 변경 파일
sql_comment = {}
create_alter = defaultdict(list)
APPLY_ORDER = [
    '2026-10-01_tqp_gds_preflight.sql', '2026-10-01_tqp_gds_map_schema.sql', '2026-10-01_tqp_gds_chain_schema.sql', '2026-10-01_tqp_gds_verify.sql',
    '2026-10-01_tqp_audit_columns.sql', '2026-10-01_tqp_gds_audit_migration.sql', '2026-10-01_tqp_gds_audit_migration_resume.sql',
    '2026-10-01_tqp_gds_audit_user_length.sql', '2026-10-01_tqp_gds_audit_verify.sql',
    '2026-10-02_tqp_gds_schema_seq.sql', '2026-10-02_tqp_gds_schema_device.sql', '2026-10-02_tqp_gds_fk_delete_rule.sql',
    '2026-10-07_tqp_gds_src_file.sql', '2026-10-07_tqp_gds_et_result.sql', '2026-10-07_tqp_gds_et_limit.sql', '2026-10-07_tqp_gds_src_file_uq_wafer.sql',
    '2026-10-12_tqp_gds_src_file_name_rule.sql', '2026-10-12_tqp_gds_src_file_uq_wafer_test.sql', '2026-10-12_tqp_gds_map_sub_device.sql',
    '2026-10-12_tqp_gds_et_limit_item.sql', '2026-10-12_tqp_gds_et_chain_map.sql', '2026-10-12_tqp_gds_et_chain_group.sql',
    '2026-10-12_tqp_gds_source_el_drop.sql',
    '2026-10-14_tqp_gds_et_para_rep_chain.sql', '2026-10-14_tqp_gds_et_limit_fk.sql', '2026-10-14_tqp_gds_et_local_reset.sql',
    '2026-10-14_tqp_gds_et_upload_registered.sql', '2026-10-14_tqp_gds_et_para_gds_chain.sql', '2026-10-14_tqp_gds_et_yms.sql',
]


def apply_key(path):
    """같은 날 파일은 실제 적용 순서(변경 이력 기준)로 정렬한다."""
    name = os.path.basename(path)
    return (name[:10], APPLY_ORDER.index(name) if name in APPLY_ORDER else -1, name)


for f in sorted(glob.glob(os.path.join(PERSIST, '*.sql')), key=apply_key):
    name = os.path.basename(f)
    text = open(f, 'rb').read().decode('utf-8-sig', 'replace')
    for m in re.finditer(r"COMMENT\s+ON\s+(?:COLUMN|TABLE)\s+(?:DMSMGR\.)?([A-Z0-9_]+)(?:\.([A-Z0-9_]+))?\s+IS\s+'((?:[^']|'')*)'", text, re.I):
        sql_comment[(m.group(1).upper(), (m.group(2) or '').upper())] = m.group(3).replace("''", "'")
    if name.startswith('2026-09-29'):
        continue  # 적용하지 않은 초안
    body = re.sub(r'--[^\n]*', '', text)
    for m in re.finditer(r'\b(CREATE|ALTER|DROP)\s+TABLE\s+(?:"?DMSMGR"?\.)?"?([A-Z0-9_]+)"?', body, re.I):
        kind, t = m.group(1).upper(), m.group(2).upper()
        if kind == 'CREATE':
            create_alter[t] = [(name, '생성')]
        elif kind == 'DROP':
            create_alter[t] = []
        elif not create_alter[t] or create_alter[t][-1][0] != name:
            create_alter[t].append((name, '변경'))

# ---------------------------------------------------------------- 메타데이터 정리
columns = defaultdict(list)
for r in cols_raw:
    columns[r[0]].append(r)
cons = defaultdict(list)
cons_by_name = {}
for r in cons_raw:
    cons[r[0]].append(r)
    cons_by_name[r[1]] = r
indexes = defaultdict(list)
for r in idx_raw:
    if r[3] == 'LOB':
        continue
    indexes[r[0]].append(r)
tbl_comment = {r[0]: (r[1] if len(r) > 1 else '') for r in tbl_raw}

FUNC_INDEX = {
    'UX_GDS_CHAIN_ENDPOINT': ('ROLE_CODE 가 INPUT / OUTPUT 인 행만: FACTORY, DEVICE_ID, CHAIN_SEQ, VERSION_NO, ROLE_CODE', 'Chain 버전당 INPUT 1개 / OUTPUT 최대 1개 (함수 기반 UNIQUE)'),
    'UX_GDS_CHAIN_CURRENT': ('IS_CURRENT = 1 인 행만: FACTORY, DEVICE_ID, CHAIN_SEQ', 'Chain 당 현재 버전 최대 1개 (함수 기반 UNIQUE)'),
    'IX_GDS_SRC_FILE_WAFER': ('FACTORY, DEVICE_ID, WAFER_ID(LOT_ID || \'-\' || 2자리 SLOT_NO)', 'Device / Wafer 로 파일 찾기'),
    'UQ_GDS_SRC_FILE_WAFER_TEST': ('FACTORY, WAFER_ID(LOT_ID || \'-\' || 2자리 SLOT_NO), TEST_PROGRAM', 'Wafer + 공정 당 파일 1개 (UQ_GDS_SRC_FILE_WAFER_TEST 제약)'),
}
INDEX_USE = {
    'IX_GDS_CHAIN_EDIT_EL': '도형 -> 편집 이력 찾기',
    'IX_GDS_CHAIN_LAYER_LY': 'Layer -> Chain 찾기 (FK 삭제 검사)',
    'IX_GDS_CHAIN_MEMBER_EL': '도형 -> 소속 Chain 찾기 (연관 분석)',
    'IX_GDS_CHAIN_RULE_L1': 'Layer -> 규칙 찾기 (FK 삭제 검사)',
    'IX_GDS_CHAIN_RULE_L2': 'Layer -> 규칙 찾기 (FK 삭제 검사)',
    'IX_GDS_ET_ITEM_CHAIN': 'Chain 이름으로 측정 항목 찾기',
    'IX_GDS_ET_ITEM_CHAIN2': '둘째 Chain 이름으로 측정 항목 찾기',
    'IX_GDS_ET_ITEM_PARA': 'Para 이름 / 조건 값으로 측정 항목 찾기',
    'IX_GDS_ET_VALUE_ITEM': '파일 + 측정 항목별 측정값 조회',
    'IX_GDS_ET_VALUE_ITEMONLY': '측정 항목 -> 측정값 찾기 (항목 삭제 검사)',
    'IX_GDS_PLACED_LAYER': 'Layer 별 도형 순차 조회',
}

def data_type(r):
    t, length, clen, used, prec, scale = r[3], r[4], r[5], r[6], r[7], r[8]
    if t in ('VARCHAR2', 'CHAR', 'NVARCHAR2'):
        return '%s(%s%s)' % (t, clen, ' CHAR' if used == 'C' else '')
    if t == 'NUMBER':
        if not prec:
            return 'NUMBER'
        return 'NUMBER(%s)' % prec if scale in ('', '0') else 'NUMBER(%s,%s)' % (prec, scale)
    return t

def description(t, c, db_comment):
    if (t, c) in DESC:
        return DESC[(t, c)]
    if db_comment and '\ufffd' not in db_comment:
        return db_comment
    if (t, c) in sql_comment:
        return sql_comment[(t, c)]
    if t.startswith('TQP_GDS_CHAIN_') and t != 'TQP_GDS_CHAIN_VER' and c in CHAIN_CHILD:
        return CHAIN_CHILD[c]
    if t == 'TQP_GDS_CHAIN_VER' and c == 'CHAIN_SEQ':
        return CHAIN_CHILD[c]
    if c in COMMON:
        return COMMON[c]
    return ''

def clean(s):
    return re.sub(r'\s+', ' ', s or '').strip()

# CHECK 에서 단일 컬럼 허용 값 뽑기
allowed = {}
for t, rows in cons.items():
    for r in rows:
        if r[2] != 'C' or r[8] == 'GENERATED NAME':
            continue
        cond = clean(r[3])
        for m in re.finditer(r"\b([A-Z_0-9]+) IN \(([^)]*)\)", cond):
            col, vals = m.group(1), [v.strip().strip("'") for v in m.group(2).split(',')]
            if re.search(r'\b%s IS NULL OR %s IN' % (col, col), cond):
                vals = vals
            allowed.setdefault((t, col), vals)
        m = re.match(r'^([A-Z_0-9]+) BETWEEN (\S+) AND (\S+)$', cond)
        if m:
            allowed.setdefault((t, m.group(1)), ['%s ~ %s' % (m.group(2), m.group(3))])
        m = re.match(r'^([A-Z_0-9]+) (>=|>) (\S+)$', cond)
        if m:
            allowed.setdefault((t, m.group(1)), ['%s %s' % (m.group(2), m.group(3))])

pk_pos, fk_parent = {}, defaultdict(list)
for t, rows in cons.items():
    for r in rows:
        cols = [x.strip() for x in r[9].split(',')] if r[9] else []
        if r[2] == 'P':
            for i, c in enumerate(cols, 1):
                pk_pos[(t, c)] = i
        elif r[2] == 'R':
            for c in cols:
                if r[5] not in fk_parent[(t, c)]:
                    fk_parent[(t, c)].append(r[5])

parents, children = defaultdict(list), defaultdict(list)
for t, rows in cons.items():
    for r in rows:
        if r[2] == 'R':
            if r[5] not in parents[t]:
                parents[t].append(r[5])
            if t not in children[r[5]]:
                children[r[5]].append(t)

def query_xml(t):
    path = os.path.join(ROOT, r'DMS\08_DBQuery\ORACLE', t + '.xml')
    return t + '.xml' if os.path.exists(path) else ''

# ---------------------------------------------------------------- 스타일
FONT = '맑은 고딕'
thin = Side(style='thin', color='A6A6A6')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
LABEL_FILL = PatternFill('solid', fgColor='DDEBF7')
PK_FILL = PatternFill('solid', fgColor='FFF2CC')
SECTION_FILL = PatternFill('solid', fgColor='BDD7EE')
WRAP = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

FORMULA_PREFIX = ('=COUNTA(', '=SUM(', '=COUNTIF(')


def keep_text(c, value):
    """'=' 로 시작하는 설명 문자열(예: '= 뒤 조건 값')이 수식으로 저장되지 않게 문자열로 고정한다."""
    if isinstance(value, str) and value.startswith('=') and not value.startswith(FORMULA_PREFIX):
        c.data_type = 's'


def text_lines(text, width):
    """병합 셀처럼 자동 높이가 안 되는 칸의 줄 수 추정 (한글 2칸, 그 밖 1칸, 열 너비 1 = 약 1.15칸)."""
    lines = 0
    for part in str(text or '').split('\n'):
        vis = sum(2 if ord(ch) > 0x2E80 else 1 for ch in part)
        lines += max(1, -(-vis // int(width * 1.15)))
    return lines


def fit(ws, row, *pairs):
    """pairs = (글자, 열 너비 합). 가장 많은 줄 수에 맞춰 행 높이를 정한다."""
    ws.row_dimensions[row].height = max(16, 14 * max(text_lines(t, w) for t, w in pairs) + 4)


def font(bold=False, size=10, color='000000', underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline)

def cell(ws, ref, value, bold=False, fill=None, align=WRAP, border=True, size=10, color='000000'):
    c = ws[ref]
    c.value = value
    keep_text(c, value)
    c.font = font(bold, size, color)
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    return c

def merge(ws, rng, value, **kw):
    first = rng.split(':')[0]
    ws.merge_cells(rng)
    c = cell(ws, first, value, **kw)
    if kw.get('border', True):
        for row in ws[rng]:
            for x in row:
                x.border = BORDER
    return c

def header_row(ws, row, titles, start_col=1):
    for i, title in enumerate(titles):
        c = ws.cell(row=row, column=start_col + i, value=title)
        c.font = font(True, 10, 'FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER

def write_row(ws, row, values, start_col=1, fill=None, centers=()):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        keep_text(c, v)
        c.font = font()
        c.alignment = CENTER if (start_col + i) in centers else WRAP
        c.border = BORDER
        if fill:
            c.fill = fill

def link(c, sheet):
    c.hyperlink = "#'%s'!A1" % sheet
    c.font = font(False, 10, '0563C1', 'single')

def setup_print(ws, landscape=True):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False

wb = Workbook()

# ================================================================ 테이블별 시트 (먼저 만들어 행 범위를 알아 둔다)
col_range = {}
table_sheets = []
for t, (group, tdesc) in TABLES.items():
    ws = wb.create_sheet(t)
    table_sheets.append(ws)
    setup_print(ws)
    widths = [5, 24, 58, 20, 7, 5, 22, 18, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    merge(ws, 'A1:I1', '테이블 명세서', bold=True, size=14, align=Alignment(vertical='center'), border=False)
    ws.row_dimensions[1].height = 26
    back = ws['I2']
    back.value = '<- 테이블 목록'
    link(back, '테이블 목록')
    back.alignment = Alignment(horizontal='right')

    pk = [r for r in cons[t] if r[2] == 'P']
    pk_text = '%s (%s)' % (pk[0][1], pk[0][9]) if pk else ''
    files = create_alter.get(t, [])
    file_text = '\n'.join('%s (%s)' % (n, k) for n, k in files)

    merge(ws, 'A3:B3', '테이블명', bold=True, fill=LABEL_FILL)
    cell(ws, 'C3', t, bold=True)
    cell(ws, 'D3', '업무 구분', bold=True, fill=LABEL_FILL)
    merge(ws, 'E3:F3', group)
    cell(ws, 'G3', '스키마', bold=True, fill=LABEL_FILL)
    merge(ws, 'H3:I3', 'DMSMGR')
    merge(ws, 'A4:B4', '테이블 설명', bold=True, fill=LABEL_FILL)
    merge(ws, 'C4:I4', tdesc)
    fit(ws, 4, (tdesc, 156))
    merge(ws, 'A5:B5', '기본 키 (PK)', bold=True, fill=LABEL_FILL)
    cell(ws, 'C5', pk_text)
    cell(ws, 'D5', '컬럼 수', bold=True, fill=LABEL_FILL)
    merge(ws, 'E5:F5', None, align=CENTER)
    cell(ws, 'G5', 'Query XML', bold=True, fill=LABEL_FILL)
    merge(ws, 'H5:I5', query_xml(t))
    merge(ws, 'A6:B6', '생성 / 변경 SQL', bold=True, fill=LABEL_FILL)
    merge(ws, 'C6:I6', file_text)
    fit(ws, 6, (file_text, 156))

    header_row(ws, 8, ['No', '컬럼명', '컬럼 설명', '데이터 타입', 'NULL', 'PK', 'FK 참조', '기본값', '허용 값'])
    row = 9
    first = row
    for r in columns[t]:
        c = r[2]
        default = defaults.get((t, c), '')
        if default.upper() == 'NULL':
            default = ''
        desc = description(t, c, r[10])
        if (t, c) in VIRTUAL:
            default = ''  # 계산식은 컬럼 설명(DB 주석)에 있다
        vals = allowed.get((t, c))
        allowed_text = ' / '.join(vals) if vals else ''
        write_row(ws, row, [int(r[1]), c, desc, data_type(r), 'N' if r[9] == 'N' else 'Y',
                            pk_pos.get((t, c), ''), ', '.join(fk_parent.get((t, c), [])), default, allowed_text],
                  fill=PK_FILL if (t, c) in pk_pos else None, centers=(1, 5, 6))
        row += 1
    last = row - 1
    col_range[t] = (first, last)
    ws['E5'].value = '=COUNTA(B%d:B%d)' % (first, last)

    # 제약 조건
    row += 1
    merge(ws, 'A%d:I%d' % (row, row), '제약 조건', bold=True, fill=SECTION_FILL)
    row += 1
    header_row(ws, row, ['No', '구분', '제약 이름', '컬럼', '', '', '내용', '', ''])
    ws.merge_cells('D%d:F%d' % (row, row))
    ws.merge_cells('G%d:I%d' % (row, row))
    row += 1
    kinds = {'P': 'PK', 'U': 'UNIQUE', 'R': 'FK', 'C': 'CHECK'}
    n = 1
    for r in cons[t]:
        if r[2] == 'C' and r[8] == 'GENERATED NAME':
            continue
        if r[2] == 'P':
            content = '기본 키'
        elif r[2] == 'U':
            content = '중복 금지'
        elif r[2] == 'R':
            parent = cons_by_name.get(r[4])
            pcols = parent[9] if parent else ''
            rule = 'CASCADE (부모 삭제 시 함께 삭제)' if r[6] == 'CASCADE' else 'NO ACTION (자식이 있으면 부모 삭제 불가)'
            content = '-> %s (%s) / 삭제 규칙 %s' % (r[5], pcols, rule)
        else:
            content = clean(r[3])
        write_row(ws, row, [n, kinds.get(r[2], r[2]), r[1], r[9] if r[2] != 'C' else '', '', '', content, '', ''], centers=(1, 2))
        ws.merge_cells('D%d:F%d' % (row, row))
        ws.merge_cells('G%d:I%d' % (row, row))
        fit(ws, row, (content, 66), (r[9], 32))
        row += 1
        n += 1

    # 인덱스
    row += 1
    merge(ws, 'A%d:I%d' % (row, row), '인덱스', bold=True, fill=SECTION_FILL)
    row += 1
    header_row(ws, row, ['No', '유일', '인덱스 이름', '컬럼 / 식', '', '', '용도', '', ''])
    ws.merge_cells('D%d:F%d' % (row, row))
    ws.merge_cells('G%d:I%d' % (row, row))
    row += 1
    for n, r in enumerate(indexes[t], 1):
        name = r[1]
        cols = r[4]
        use = INDEX_USE.get(name, '')
        if name in FUNC_INDEX:
            cols, use = FUNC_INDEX[name]
        elif r[5] == 'P':
            use = '기본 키'
        elif r[5] == 'U':
            use = '중복 금지 (UNIQUE 제약)'
        write_row(ws, row, [n, 'UNIQUE' if r[2] == 'UNIQUE' else '-', name, cols, '', '', use, '', ''], centers=(1, 2))
        ws.merge_cells('D%d:F%d' % (row, row))
        ws.merge_cells('G%d:I%d' % (row, row))
        fit(ws, row, (cols, 32), (use, 66))
        row += 1
    ws.freeze_panes = 'A9'

# ================================================================ 테이블 목록
ws = wb.create_sheet('테이블 목록', 0)
setup_print(ws)
for i, w in enumerate([5, 11, 24, 52, 36, 8, 22, 26, 26, 54], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
merge(ws, 'A1:J1', 'POC 추가 테이블 목록', bold=True, size=14, align=Alignment(vertical='center'), border=False)
ws.row_dimensions[1].height = 26
header_row(ws, 3, ['No', '업무 구분', '테이블명', '테이블 설명', '기본 키 (PK)', '컬럼 수', '상위 테이블 (FK)', '하위 테이블', 'Query XML', '생성 / 변경 SQL'])
row = 4
for n, (t, (group, tdesc)) in enumerate(TABLES.items(), 1):
    pk = [r for r in cons[t] if r[2] == 'P']
    first, last = col_range[t]
    files = create_alter.get(t, [])
    write_row(ws, row, [n, group, t, tdesc, pk[0][9] if pk else '', '=COUNTA(\'%s\'!B%d:B%d)' % (t, first, last),
                        ', '.join(parents[t]), ', '.join(children[t]), query_xml(t),
                        '\n'.join('%s (%s)' % (f, k) for f, k in files)], centers=(1, 2, 6))
    link(ws.cell(row=row, column=3), t)
    row += 1
list_first, list_last = 4, row - 1
cell(ws, 'E%d' % row, '합계', bold=True, fill=LABEL_FILL, align=CENTER)
cell(ws, 'F%d' % row, '=SUM(F%d:F%d)' % (list_first, list_last), bold=True, fill=LABEL_FILL, align=CENTER)
ws.freeze_panes = 'A4'

# ================================================================ 관계 (FK)
ws = wb.create_sheet('관계(FK)')
setup_print(ws)
for i, w in enumerate([5, 24, 30, 40, 22, 40, 14, 44], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
merge(ws, 'A1:H1', '테이블 관계 (FK)', bold=True, size=14, align=Alignment(vertical='center'), border=False)
ws.row_dimensions[1].height = 26
tree = [
    'TQP_GDS_MAP (Factory + Device 당 1건)',
    '  +- TQP_GDS_LAYER',
    '  +- TQP_GDS_PLACED_EL ........................ TQP_GDS_LAYER 참조',
    '  +- TQP_GDS_IMPORT_JOB',
    '  |     +- TQP_GDS_IMPORT_BATCH',
    '  +- TQP_GDS_CHAIN',
    '        +- TQP_GDS_CHAIN_VER',
    '              +- TQP_GDS_CHAIN_MEMBER .......... TQP_GDS_PLACED_EL 참조',
    '              +- TQP_GDS_CHAIN_LAYER ........... TQP_GDS_LAYER 참조',
    '              +- TQP_GDS_CHAIN_RULE ............ TQP_GDS_LAYER 참조 (Layer 1 / Layer 2)',
    '              +- TQP_GDS_CHAIN_EDIT ............ TQP_GDS_PLACED_EL 참조',
    '',
    'TQP_GDS_SRC_FILE (E-Test 파일 1개)',
    '  +- TQP_GDS_ET_DIE',
    '        +- TQP_GDS_ET_VALUE .................... TQP_GDS_ET_ITEM 참조',
    'TQP_GDS_ET_ITEM (등록 Parameter)',
    '  +- TQP_GDS_ET_LIMIT',
    '',
    'E-Test 와 GDS 는 FK 없이 이름으로 연결한다: SRC_FILE / ET_ITEM 의 DEVICE_ID = TQP_GDS_MAP.DEVICE_ID, ET_ITEM.GDS_CHAIN_NAME = TQP_GDS_CHAIN_VER.CHAIN_NAME',
]
row = 3
for line in tree:
    c = ws.cell(row=row, column=2, value=line)
    c.font = Font(name='Consolas', size=10)
    row += 1
row += 1
header_row(ws, row, ['No', '자식 테이블', 'FK 이름', '자식 컬럼', '부모 테이블', '부모 키', '삭제 규칙', '비고'])
row += 1
n = 1
for t in TABLES:
    for r in cons[t]:
        if r[2] != 'R':
            continue
        parent = cons_by_name.get(r[4])
        note = '부모를 지우면 함께 지워진다' if r[6] == 'CASCADE' else '자식이 있으면 부모를 지울 수 없다 (자식을 먼저 지워야 함)'
        write_row(ws, row, [n, t, r[1], r[9], r[5], parent[9] if parent else '', r[6], note], centers=(1, 7))
        row += 1
        n += 1
ws.freeze_panes = 'A%d' % (len(tree) + 5)

# ================================================================ 시퀀스
ws = wb.create_sheet('시퀀스')
setup_print(ws)
for i, w in enumerate([5, 20, 36, 8, 8, 8, 34, 44], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
merge(ws, 'A1:H1', '시퀀스', bold=True, size=14, align=Alignment(vertical='center'), border=False)
ws.row_dimensions[1].height = 26
header_row(ws, 3, ['No', '시퀀스', '사용 테이블.컬럼', '시작', '증가', 'CACHE', '생성 SQL', '비고'])
SEQS = [
    ('IMPORT_JOB_SEQ', 'TQP_GDS_IMPORT_JOB.IMPORT_JOB_SEQ', '2026-10-02_tqp_gds_schema_seq.sql', ''),
    ('CHAIN_SEQ', 'TQP_GDS_CHAIN.CHAIN_SEQ', '2026-10-02_tqp_gds_schema_seq.sql', ''),
    ('BATCH_EDIT_SEQ', 'TQP_GDS_CHAIN_EDIT.BATCH_EDIT_SEQ', '2026-10-02_tqp_gds_schema_seq.sql', '현재 미사용 (일괄 편집 묶음용)'),
    ('SRC_FILE_SEQ', 'TQP_GDS_SRC_FILE.SRC_FILE_SEQ', '2026-10-07_tqp_gds_src_file.sql', ''),
    ('ET_ITEM_SEQ', 'TQP_GDS_ET_ITEM.ET_ITEM_SEQ', '2026-10-07_tqp_gds_et_result.sql', ''),
]
for n, (s, use, f, note) in enumerate(SEQS, 1):
    write_row(ws, 3 + n, [n, s, use, 1, 1, 20, f, note], centers=(1, 4, 5, 6))
r = 4 + len(SEQS) + 1
merge(ws, 'A%d:H%d' % (r, r), '번호는 SELECT <시퀀스>.NEXTVAL FROM DUAL 로 받는다 (기존 DMS 의 STEP_SEQ 등과 같은 방식). INSERT 가 실패해도 번호가 소모되어 빈 번호가 생길 수 있다.', border=False)

# ================================================================ 변경 이력
ws = wb.create_sheet('변경 이력')
setup_print(ws)
for i, w in enumerate([5, 12, 46, 46, 72], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
merge(ws, 'A1:E1', '변경 이력', bold=True, size=14, align=Alignment(vertical='center'), border=False)
ws.row_dimensions[1].height = 26

row = 3
merge(ws, 'A%d:E%d' % (row, row), '1. 개발 중 만들었다가 삭제한 테이블 (현재 DB 에 없음)', bold=True, fill=SECTION_FILL)
row += 1
header_row(ws, row, ['No', '삭제일', '테이블', '생성 -> 삭제 SQL', '삭제 사유 / 대체'])
DROPPED = [
    ('2026-10-02', 'TQP_GDS_MAP_REV', '2026-10-01_tqp_gds_map_schema.sql\n-> 2026-10-02_tqp_gds_schema_device.sql', 'Map 개정본(Revision) 관리를 하지 않기로 함. Factory + Device 1:1 헤더 TQP_GDS_MAP 으로 대체'),
    ('2026-10-02', 'TQP_GDS_MAP_ACTIVE', '2026-10-01_tqp_gds_map_schema.sql\n-> 2026-10-02_tqp_gds_schema_device.sql', '현재 사용 개정본 선택 테이블. 개정본 관리 제거로 불필요'),
    ('2026-10-12', 'TQP_GDS_SOURCE_EL', '2026-10-01_tqp_gds_map_schema.sql\n-> 2026-10-12_tqp_gds_source_el_drop.sql', 'GDS 원본 행. 저장만 하고 조회하는 화면 / 서비스가 없어 삭제. 도면 / Chain / Label 검색은 TQP_GDS_PLACED_EL 만 사용'),
    ('2026-10-12', 'TQP_GDS_ET_CHAIN_MAP', '2026-10-12_tqp_gds_et_chain_map.sql\n-> 2026-10-12_tqp_gds_et_chain_group.sql', 'E-Test Chain -> 대표 Chain 연결. 같은 날 묶음 번호 방식(TQP_GDS_ET_CHAIN_GROUP)으로 대체'),
    ('2026-10-14', 'TQP_GDS_ET_CHAIN_GROUP', '2026-10-12_tqp_gds_et_chain_group.sql\n-> 2026-10-14_tqp_gds_et_limit_fk.sql', 'E-Test Chain 묶음 (ET_CHAIN_GROUP_SEQ 시퀀스 포함). 대표 Chain(TQP_GDS_ET_REP_CHAIN) 방식으로 대체'),
    ('2026-10-14', 'TQP_GDS_ET_REP_CHAIN', '2026-10-14_tqp_gds_et_para_rep_chain.sql\n-> 2026-10-14_tqp_gds_et_para_gds_chain.sql', '대표 Chain. Parameter 별 GDS Chain 이름(TQP_GDS_ET_ITEM.GDS_CHAIN_NAME) 직접 입력으로 대체'),
    ('2026-10-14', 'TQP_GDS_ET_CHAIN_LINK', '2026-10-14_tqp_gds_et_para_rep_chain.sql\n-> 2026-10-14_tqp_gds_et_para_gds_chain.sql', '대표 Chain 과 E-Test Chain 연결. 위와 같은 이유로 삭제'),
    ('2026-10-15', 'TQP_GDS_MAP_SUB_DEVICE', '2026-10-12_tqp_gds_map_sub_device.sql\n-> 2026-10-15_tqp_gds_map_sub_device_drop.sql', '대표 Device / Sub Device 연결. GDS Map 복사(같은 도면을 다른 Device 이름으로 저장)로 대체'),
]
row += 1
for n, d in enumerate(DROPPED, 1):
    write_row(ws, row, [n] + list(d), centers=(1, 2))
    row += 1

row += 1
merge(ws, 'A%d:E%d' % (row, row), '2. POC 에서 바꾼 기존 테이블 (구조 변경)', bold=True, fill=SECTION_FILL)
row += 1
header_row(ws, row, ['No', '변경일', '테이블.컬럼', 'SQL', '변경 내용 / 사유'])
EXISTING = [
    ('2026-10-07', 'TQD_SETUP.SETUP_ID', 'DMS/Assistant/DB/sql/2026-10-07_tqd_setup_setup_id_128.sql', 'VARCHAR2(40) -> VARCHAR2(128 BYTE). 고객 KLARF SetupID 가 46자라 KLARF 업로드 저장이 실패(ORA-12899)해서 확대. 고객 DB 에도 같은 DDL 필요'),
    ('2026-10-14', 'TQP_MAPDEF.FACTORY', 'DMS/Assistant/DB/sql/2026-10-14_tqp_mapdef_factory_sck.sql', "기본값 'FAB2' -> 'SCK' (Factory 통합). Map 저장 쿼리가 FACTORY 를 넣지 않아 기본값으로 저장되기 때문. 고객 DB 도 기본값 확인 필요"),
]
row += 1
for n, d in enumerate(EXISTING, 1):
    write_row(ws, row, [n] + list(d), centers=(1, 2))
    row += 1

row += 1
merge(ws, 'A%d:E%d' % (row, row), '3. DDL 스크립트 이력 (DMS/06_CoreLibSource/GdsMapControl/Persistence, 로컬 개발 DB 적용 순서)', bold=True, fill=SECTION_FILL)
row += 1
header_row(ws, row, ['No', '날짜', 'SQL 파일', '대상 테이블', '내용'])
HISTORY = [
    ('2026-10-01_tqp_gds_map_schema.sql / 2026-10-01_tqp_gds_chain_schema.sql', 'GDS 13개', '최초 GDS Map / Chain 테이블 13개 생성 (GUID 키, Map 개정본 관리). 이후 재생성되어 이력으로만 남음'),
    ('2026-10-01_tqp_gds_audit_migration*.sql / audit_user_length.sql', 'GDS 13개', '감사 컬럼 이름 / 형식을 기존 TQP 테이블과 맞춤 (CREATE_TIME / UPDATE_TIME DATE, 사용자 30 BYTE)'),
    ('2026-10-02_tqp_gds_schema_seq.sql', 'GDS 13개', '키를 GUID 에서 시퀀스(NUMBER)로 변경. 시퀀스 생성'),
    ('2026-10-02_tqp_gds_schema_device.sql', 'GDS 12개', '개정본 관리 제거. 키를 FACTORY + DEVICE_ID 로 바꿔 재생성, TQP_GDS_MAP 신규'),
    ('2026-10-02_tqp_gds_fk_delete_rule.sql', 'GDS 큰 테이블 사이 FK', '큰 테이블 사이 FK 7개의 CASCADE 제거 + 인덱스 3개 (Device 1건 삭제가 10분을 넘겨 변경)'),
    ('2026-10-06_tqp_gds_map_sha_nullable.sql', 'TQP_GDS_MAP', 'GDS_SHA256 NULL 허용 (화면에서 파일 해시 계산 제거)'),
    ('2026-10-06_tqp_gds_map_view_side.sql', 'TQP_GDS_MAP', 'VIEW_SIDE 추가 (Front / Back side 저장 기준)'),
    ('2026-10-07_tqp_gds_src_file.sql', 'TQP_GDS_SRC_FILE', 'E-Test 결과 파일 테이블 + SRC_FILE_SEQ 생성'),
    ('2026-10-07_tqp_gds_et_result.sql', 'ET_ITEM / ET_DIE / ET_VALUE / SRC_FILE', 'E-Test 측정 항목 / Die / 측정값 테이블 + ET_ITEM_SEQ 생성, 파일 테이블 건수 / 상태 컬럼 추가'),
    ('2026-10-07_tqp_gds_et_limit.sql', 'TQP_GDS_ET_LIMIT', 'E-Test Limit 최초 생성 (Para 이름 기준). 10-12 재생성'),
    ('2026-10-07_tqp_gds_src_file_uq_wafer.sql', 'TQP_GDS_SRC_FILE', '파일 중복 기준을 FACTORY + WAFER_ID 로 변경'),
    ('2026-10-08_tqp_gds_map_drop_sha.sql', 'TQP_GDS_MAP', 'GDS_SHA256 컬럼 삭제'),
    ('2026-10-12_tqp_gds_src_file_name_rule.sql', 'TQP_GDS_SRC_FILE', '파일명 8칸 규칙 컬럼 추가 (CUSTOMER / FILE_DEVICE / WAFER_TYPE / WAFER_REV / LOT_WAFER_NO / FAB_NAME / Pass / Fail 건수 / PARSE_RULE_VER / VIEW_SIDE)'),
    ('2026-10-12_tqp_gds_src_file_uq_wafer_test.sql', 'TQP_GDS_SRC_FILE', '파일 중복 기준을 FACTORY + WAFER_ID + TEST_PROGRAM(공정) 으로 변경'),
    ('2026-10-12_tqp_gds_map_sub_device.sql', 'TQP_GDS_MAP_SUB_DEVICE', '대표 Device / Sub Device 테이블 생성 (10-15 삭제)'),
    ('2026-10-12_tqp_gds_et_limit_item.sql', 'TQP_GDS_ET_LIMIT', 'Spec 을 Device + 공정 + Parameter(헤더 전체) 키로 재생성'),
    ('2026-10-12_tqp_gds_et_chain_map.sql', 'TQP_GDS_ET_CHAIN_MAP', 'E-Test Chain -> 대표 Chain 연결 생성 (같은 날 대체됨, 다시 실행하지 않음)'),
    ('2026-10-12_tqp_gds_et_chain_group.sql', 'TQP_GDS_ET_CHAIN_GROUP', 'Chain 묶음 테이블 + 시퀀스 생성, ET_CHAIN_MAP 삭제'),
    ('2026-10-12_tqp_gds_source_el_drop.sql', 'SOURCE_EL / MAP / IMPORT_JOB / IMPORT_BATCH / PLACED_EL', 'TQP_GDS_SOURCE_EL 삭제, 원본 건수 / 원본 Payload 버전 컬럼과 SOURCE 배치 종류 정리'),
    ('2026-10-14_tqp_gds_et_para_rep_chain.sql', 'TQP_GDS_ET_ITEM / ET_REP_CHAIN / ET_CHAIN_LINK', 'ET_ITEM 을 Parameter 마스터로 쓰는 컬럼 추가, 대표 Chain 테이블 2개 생성'),
    ('2026-10-14_tqp_gds_et_limit_fk.sql', 'TQP_GDS_ET_LIMIT / ET_CHAIN_GROUP', 'Spec -> 등록 Parameter FK (CASCADE), Chain 묶음 테이블 삭제'),
    ('2026-10-14_tqp_gds_et_local_reset.sql', 'E-Test 데이터', '로컬 DB 전용 E-Test 데이터 정리 (구조 변경 없음, 고객 DB 실행 금지)'),
    ('2026-10-14_tqp_gds_et_upload_registered.sql', 'SRC_FILE / ET_DIE / ET_VALUE / ET_ITEM', '업로드 재설계: 헤더 / 등록 / 누락 건수, Spec 판정, Map 좌표 컬럼 추가, ET_ITEM.TEST_PROGRAM NOT NULL'),
    ('2026-10-14_tqp_gds_et_para_gds_chain.sql', 'TQP_GDS_ET_ITEM / ET_REP_CHAIN / ET_CHAIN_LINK', 'ET_ITEM.GDS_CHAIN_NAME 추가, 대표 Chain 테이블 2개 삭제'),
    ('2026-10-14_tqp_gds_et_yms.sql', 'TQP_GDS_SRC_FILE', 'YMS 반영 상태 컬럼 추가 (YMS_STATUS / PROGRAM / LOT_SEQ / WAFER_SEQ / TIME / MESSAGE)'),
    ('2026-10-15_tqp_gds_map_sub_device_drop.sql', 'TQP_GDS_MAP_SUB_DEVICE', 'Sub Device 테이블 삭제 (GDS Map 복사로 대체)'),
    ('2026-10-15_tqp_gds_et_comment_refresh.sql', 'TQP_GDS_ET_ITEM / TQP_GDS_SRC_FILE', 'E-Test 주석 19개를 현재 동작 기준으로 갱신 (구조 변경 없음). 문자가 깨져 저장된 YMS 컬럼 주석 6개 다시 넣음'),
]
row += 1
for n, (f, target, content) in enumerate(HISTORY, 1):
    write_row(ws, row, [n, f[:10], f, target, content], centers=(1, 2))
    row += 1

# ================================================================ 개요 (맨 앞)
ws = wb['Sheet']
ws.title = '개요'
wb.move_sheet(ws, offset=-wb.index(ws))
setup_print(ws, landscape=False)
for i, w in enumerate([4, 22, 90], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
merge(ws, 'B2:C2', 'SCK 분석시스템 구축 POC - 추가 테이블 명세서', bold=True, size=16, align=Alignment(vertical='center'), border=False)
ws.row_dimensions[2].height = 32
info = [
    ('대상 시스템', 'DACrux DMS (Defect / E-Test 분석, GDS Map / Chain)'),
    ('기준 DB', '로컬 개발 DB (Oracle 19c, 스키마 DMSMGR, 문자셋 AL32UTF8)'),
    ('조회 일자', TODAY),
    ('작성 근거', 'DB 메타데이터 (컬럼 / 주석 / 제약 / 인덱스 / 시퀀스) + DDL 스크립트 (DMS/06_CoreLibSource/GdsMapControl/Persistence/*.sql) + 설계 문서 (문서/설계/02_GDS, 04_ETest)'),
    ('대상 범위', 'POC 에서 새로 만든 테이블 %d개 / 시퀀스 %d개. 기존 테이블 구조 변경 %d건과 개발 중 삭제한 테이블 %d개는 [변경 이력] 시트' % (len(TABLES), len(SEQS), len(EXISTING), len(DROPPED))),
]
row = 4
for k, v in info:
    cell(ws, 'B%d' % row, k, bold=True, fill=LABEL_FILL)
    cell(ws, 'C%d' % row, v)
    row += 1

row += 1
cell(ws, 'B%d' % row, '업무 구분', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
cell(ws, 'C%d' % row, '테이블 수', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
row += 1
grp_first = row
for g, note in [('GDS Map', 'GDS 도면 저장 (헤더 / Layer / 배치 도형 / 저장 작업)'), ('GDS Chain', 'Chain 설정 / 포함 도형 / 탐색 Layer / 규칙 / 편집 이력'), ('E-Test', 'E-Test 결과 파일 / 측정 항목 / Die / 측정값 / Spec')]:
    cell(ws, 'B%d' % row, g, bold=True)
    c = cell(ws, 'C%d' % row, "=COUNTIF('테이블 목록'!B%d:B%d,B%d)" % (list_first, list_last, row), align=Alignment(horizontal='left', vertical='center'))
    c.number_format = '0"개  - %s"' % note
    row += 1
cell(ws, 'B%d' % row, '합계', bold=True, fill=LABEL_FILL)
c = cell(ws, 'C%d' % row, '=SUM(C%d:C%d)' % (grp_first, row - 1), bold=True, fill=LABEL_FILL, align=Alignment(horizontal='left', vertical='center'))
c.number_format = '0"개"'
row += 2

cell(ws, 'B%d' % row, '시트', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
cell(ws, 'C%d' % row, '내용', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
row += 1
for sheet, v in [('테이블 목록', '테이블 16개 요약 (테이블명을 누르면 해당 시트로 이동)'),
                 ('TQP_GDS_*', '테이블별 명세: 테이블 정보 / 컬럼 / 제약 조건 / 인덱스'),
                 ('관계(FK)', '테이블 관계도와 FK 목록 (삭제 규칙 포함)'),
                 ('시퀀스', '키 번호용 시퀀스 5개'),
                 ('변경 이력', '삭제한 테이블 / 바꾼 기존 테이블 / DDL 스크립트 이력')]:
    c = cell(ws, 'B%d' % row, sheet, bold=True)
    if sheet in wb.sheetnames:
        link(c, sheet)
        c.font = font(True, 10, '0563C1', 'single')
    cell(ws, 'C%d' % row, v)
    row += 1
row += 1

cell(ws, 'B%d' % row, '표기 기준', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
cell(ws, 'C%d' % row, '설명', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
row += 1
for k, v in [('데이터 타입', 'VARCHAR2(n CHAR) = 글자 수 기준, VARCHAR2(n) = byte 기준. BINARY_DOUBLE = 8 byte 실수 (좌표 / 측정값)'),
             ('NULL', 'N = NOT NULL (필수), Y = NULL 허용'),
             ('PK', '기본 키 컬럼과 키 안 순서. 노란 행 = 기본 키 컬럼'),
             ('FK 참조', '이 컬럼이 들어 있는 FK 의 부모 테이블. 상세 컬럼 / 삭제 규칙은 테이블 시트의 제약 조건 또는 [관계(FK)] 시트'),
             ('허용 값', 'CHECK 제약으로 정한 값. 여러 컬럼을 함께 검사하는 조건은 제약 조건 표에 원문으로 적음'),
             ('감사 컬럼', 'CREATE_USER / CREATE_TIME / UPDATE_USER / UPDATE_TIME. CREATE_TIME 은 기본값 SYSDATE, UPDATE_TIME 은 Service 가 저장할 때 직접 기록')]:
    cell(ws, 'B%d' % row, k, bold=True)
    cell(ws, 'C%d' % row, v)
    row += 1
row += 1

cell(ws, 'B%d' % row, '참고', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
cell(ws, 'C%d' % row, '내용', bold=True, fill=HEAD_FILL, color='FFFFFF', align=CENTER)
row += 1
for k, v in [('고객 DB', '이 명세는 로컬 개발 DB 기준이다. 고객 DB 의 적용 상태는 확인하지 않았다.'),
             ('Chain 버전', 'Chain 은 현재 버전 이력 없이 VERSION_NO = 1 로 덮어써 저장한다 (구조는 버전 이력을 담을 수 있게 되어 있음).'),
             ('E-Test 와 GDS 연결', 'FK 없이 이름으로 연결한다 (DEVICE_ID, GDS Chain 이름).'),
             ('DB 주석', 'GDS Map / Chain 테이블은 DB 주석이 없어 설계 문서 기준으로 설명을 적었다. 그 밖의 컬럼은 DB 주석을 쓴다.'),
             ('제외', 'TD_<Program> 테이블 (예: TD_OMM_TSVET) 은 기존 YMS 적재 방식이 Program 별로 만드는 데이터 테이블이라 대상에서 뺐다.')]:
    cell(ws, 'B%d' % row, k, bold=True)
    cell(ws, 'C%d' % row, v)
    row += 1

wb.active = 0
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('저장:', OUT)
print('테이블 %d개 / 컬럼 %d개' % (len(TABLES), sum(len(columns[t]) for t in TABLES)))
# 확인: 목록(TABLES)에 없는 POC 테이블 / 설명이 빈 컬럼
extra = sorted(set(tbl_comment) - set(TABLES))
if extra:
    print('경고: TABLES 에 없는 TQP_GDS 테이블 ->', ', '.join(extra))
for t in TABLES:
    missing = [r[2] for r in columns[t] if not description(t, r[2], r[10])]
    if not columns[t]:
        print('경고: DB 에 없는 테이블 ->', t)
    if missing:
        print('경고: 설명이 빈 컬럼', t, missing)
